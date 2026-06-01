"""
Tests for the file browser preview views (preview_file, inline_file).
"""

from io import BytesIO

from django.core.files.base import ContentFile
from django.template.loader import render_to_string
from django.urls import reverse

import pytest
from chat_next.models import Chat, ChatFile, Message
from docx import Document as WordDocument

from librarian.models import Document, SavedFile


@pytest.mark.django_db
def test_preview_file_markdown(client, all_apps_user):
    """preview_file returns markdown-text div for .md files."""
    user = all_apps_user()
    client.force_login(user)

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="", is_bot=True)
    sf = SavedFile.objects.create(content_type="text/markdown")
    sf.file.save("output.md", ContentFile(b"# Hello World\n\nSome **bold** text."))
    cf = ChatFile.objects.create(message=msg, filename="output.md", saved_file=sf)

    url = reverse("chat_next:preview_file", args=[cf.id])
    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode()
    assert 'class="markdown-text"' in content
    assert 'data-copyable="true"' in content
    assert "data-md=" in content
    assert "Hello World" in content


@pytest.mark.django_db
def test_preview_file_plain_text(client, all_apps_user):
    """preview_file returns <pre><code> for text files."""
    user = all_apps_user()
    client.force_login(user)

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="", is_bot=True)
    sf = SavedFile.objects.create(content_type="text/plain")
    sf.file.save("notes.txt", ContentFile(b"Line 1\nLine 2\nLine 3"))
    cf = ChatFile.objects.create(message=msg, filename="notes.txt", saved_file=sf)

    url = reverse("chat_next:preview_file", args=[cf.id])
    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode()
    assert "<pre" in content
    assert 'data-copyable="true"' in content
    assert "Line 1" in content


@pytest.mark.django_db
def test_preview_file_image(client, all_apps_user):
    """preview_file returns <img> tag for image files."""
    user = all_apps_user()
    client.force_login(user)

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="", is_bot=True)
    sf = SavedFile.objects.create(content_type="image/png")
    sf.file.save("chart.png", ContentFile(b"\x89PNG fake image content"))
    cf = ChatFile.objects.create(message=msg, filename="chart.png", saved_file=sf)

    url = reverse("chat_next:preview_file", args=[cf.id])
    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode()
    assert "<img" in content
    assert 'data-copyable="true"' in content
    assert "chart.png" in content


@pytest.mark.django_db
def test_preview_file_pdf(client, all_apps_user):
    """preview_file returns <iframe> for PDF files."""
    user = all_apps_user()
    client.force_login(user)

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="", is_bot=True)
    sf = SavedFile.objects.create(content_type="application/pdf")
    sf.file.save("report.pdf", ContentFile(b"%PDF-1.4 fake"))
    cf = ChatFile.objects.create(message=msg, filename="report.pdf", saved_file=sf)

    url = reverse("chat_next:preview_file", args=[cf.id])
    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode()
    assert "<iframe" in content
    assert 'data-copyable="false"' in content
    assert "inline" in content  # inline URL is referenced


@pytest.mark.django_db
def test_preview_file_docx(client, all_apps_user):
    """preview_file returns a docx-preview render target for DOCX files."""
    user = all_apps_user()
    client.force_login(user)

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="", is_bot=True)

    doc = WordDocument()
    doc.add_heading("DOCX preview", level=1)
    doc.add_paragraph("This document should render inside the preview widget.")
    buffer = BytesIO()
    doc.save(buffer)

    sf = SavedFile.objects.create(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    sf.file.save("brief.docx", ContentFile(buffer.getvalue()))
    cf = ChatFile.objects.create(message=msg, filename="brief.docx", saved_file=sf)

    url = reverse("chat_next:preview_file", args=[cf.id])
    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode()
    assert 'data-preview-type="docx"' in content
    assert 'data-copyable="false"' in content
    assert "preview-docx-wrapper" in content
    assert "preview-docx-stage" in content
    assert "Loading DOCX preview" in content
    assert "inline" in content


@pytest.mark.django_db
def test_preview_file_unsupported(client, all_apps_user):
    """preview_file returns download fallback for unsupported types."""
    user = all_apps_user()
    client.force_login(user)

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="", is_bot=True)
    sf = SavedFile.objects.create(content_type="application/zip")
    sf.file.save("archive.zip", ContentFile(b"PK fake zip"))
    cf = ChatFile.objects.create(message=msg, filename="archive.zip", saved_file=sf)

    url = reverse("chat_next:preview_file", args=[cf.id])
    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode()
    assert "preview-unsupported" in content
    assert 'data-copyable="false"' in content
    assert "Download" in content


@pytest.mark.django_db
def test_inline_file_serves_inline(client, all_apps_user):
    """inline_file returns file with Content-Disposition: inline."""
    user = all_apps_user()
    client.force_login(user)

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="", is_bot=True)
    sf = SavedFile.objects.create(content_type="image/png")
    sf.file.save("chart.png", ContentFile(b"\x89PNG fake image"))
    cf = ChatFile.objects.create(message=msg, filename="chart.png", saved_file=sf)

    url = reverse("chat_next:inline_file", args=[cf.id])
    response = client.get(url)
    assert response.status_code == 200
    # FileResponse with as_attachment=False should have inline disposition
    content_disposition = response.get("Content-Disposition", "")
    assert "attachment" not in content_disposition


@pytest.mark.django_db
def test_preview_file_wrong_user(client, all_apps_user):
    """preview_file denies access to files belonging to other users."""
    owner = all_apps_user("owner")
    other = all_apps_user("other")

    chat = Chat.objects.create(user=owner)
    msg = Message.objects.create(chat=chat, text="", is_bot=True)
    sf = SavedFile.objects.create(content_type="text/plain")
    sf.file.save("secret.txt", ContentFile(b"secret"))
    cf = ChatFile.objects.create(message=msg, filename="secret.txt", saved_file=sf)

    client.force_login(other)
    url = reverse("chat_next:preview_file", args=[cf.id])
    response = client.get(url)
    assert response.status_code != 200


@pytest.mark.django_db
def test_preview_file_404(client, all_apps_user):
    """preview_file returns 404 for non-existing file."""
    user = all_apps_user()
    client.force_login(user)

    url = reverse("chat_next:preview_file", args=[99999])
    response = client.get(url)
    assert response.status_code == 404


@pytest.mark.django_db
def test_message_files_renders_file_browser(client, all_apps_user):
    """Messages with files render the file-browser component."""
    user = all_apps_user()
    client.force_login(user)

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="Here are your files", is_bot=True)
    sf = SavedFile.objects.create(content_type="text/markdown")
    sf.file.save("output.md", ContentFile(b"# Test"))
    ChatFile.objects.create(message=msg, filename="output.md", saved_file=sf)

    # Load the chat page
    url = reverse("chat_next:chat", args=[chat.id])
    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode()
    assert "file-browser" in content
    assert "file-status-card" in content
    assert "bi-eye" in content
    assert "bi-download" in content
    assert "preview-copy-btn" in content
    assert "preview-copy-dropdown" in content
    assert "preview-copy-menu-toggle" in content
    assert "Copy as markdown" in content
    assert "Copy as rich text" in content


@pytest.mark.django_db
def test_bot_message_file_browser_hides_library_status_summary(all_apps_user):
    """Bot output files should not show library status messaging."""
    user = all_apps_user()

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="Here are your files", is_bot=True)
    sf = SavedFile.objects.create(content_type="text/plain")
    sf.file.save("output.txt", ContentFile(b"done"))
    ChatFile.objects.create(message=msg, filename="output.txt", saved_file=sf)

    html = render_to_string(
        "chat_next/components/file_browser.html",
        {
            "message": Message.objects.prefetch_related("files").get(id=msg.id),
            "status_summary": {
                "processing": False,
                "success_count": 2,
                "error_count": 0,
                "stopped_count": 0,
            },
            "data_source_id": 123,
        },
    )

    assert "file-browser" in html
    assert "upload-status-summary" not in html
    assert "completed" not in html
    assert "See details" not in html


@pytest.mark.django_db
def test_bot_message_child_files_render_flat_without_child_indent(all_apps_user):
    user = all_apps_user()

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="Here are your files", is_bot=True)
    parent_saved = SavedFile.objects.create(content_type="application/pdf")
    parent_saved.file.save("source.pdf", ContentFile(b"%PDF-1.4 fake"))
    parent_doc = Document.objects.create(
        data_source=chat.data_source,
        saved_file=parent_saved,
        filename="source.pdf",
    )

    child_saved = SavedFile.objects.create(content_type="text/markdown")
    child_saved.file.save("source__chunk_001.md", ContentFile(b"# Chunk"))
    child_doc = Document.objects.create(
        data_source=chat.data_source,
        saved_file=child_saved,
        filename="source__chunk_001.md",
        parent_document=parent_doc,
    )
    ChatFile.objects.create(
        message=msg,
        filename="source__chunk_001.md",
        saved_file=child_saved,
        document=child_doc,
    )

    html = render_to_string(
        "chat_next/components/file_browser.html",
        {
            "message": Message.objects.prefetch_related(
                "files__document__parent_document"
            ).get(id=msg.id),
        },
    )

    assert "file-status-card" in html
    assert "ms-3" not in html
    assert "bi-arrow-return-right" not in html


@pytest.mark.django_db
def test_file_browser_embeds_chat_next_document_modal_url(all_apps_user):
    """Preview metadata should include the chat_next librarian modal URL for document icons."""
    user = all_apps_user()

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="Has uploads", is_bot=False)
    sf = SavedFile.objects.create(content_type="text/plain")
    sf.file.save("output.txt", ContentFile(b"done"))
    chat_file = ChatFile.objects.create(
        message=msg, filename="output.txt", saved_file=sf
    )

    from librarian.models import DataSource, Document, Library

    library = Library.objects.create(name=" ", created_by=user)
    data_source = DataSource.objects.create(library=library, name="Chat files")
    document = Document.objects.create(
        data_source=data_source,
        filename="output.txt",
        status="PAUSED",
        saved_file=sf,
    )
    document.chat_next_messages.add(msg)
    chat_file.document = document
    chat_file.save(update_fields=["document"])

    html = render_to_string(
        "chat_next/components/file_browser.html",
        {
            "message": Message.objects.prefetch_related("files__document").get(
                id=msg.id
            ),
        },
    )

    assert '"documentModalUrl":"' in html
    assert (
        reverse("chat_next:modal_librarian_document", args=[chat.id, document.id])
        in html
    )


@pytest.mark.django_db
def test_user_message_file_browser_uses_adding_to_library_summary(all_apps_user):
    """User upload summaries should use the simplified library wording."""
    user = all_apps_user()

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="Has uploads", is_bot=False)
    sf = SavedFile.objects.create(content_type="text/plain")
    sf.file.save("output.txt", ContentFile(b"done"))
    ChatFile.objects.create(message=msg, filename="output.txt", saved_file=sf)

    html = render_to_string(
        "chat_next/components/file_browser.html",
        {
            "message": Message.objects.prefetch_related("files").get(id=msg.id),
            "status_summary": {
                "processing": True,
                "processed": 1,
                "total": 2,
                "paused_count": 0,
                "success_count": 0,
                "error_count": 0,
                "stopped_count": 0,
            },
            "data_source_id": 123,
        },
    )

    assert "Adding to library... (1/2)" in html
    assert "Q&amp;A library" not in html


@pytest.mark.django_db
def test_preview_file_html(client, all_apps_user):
    """preview_file returns sandboxed iframe for HTML files."""
    user = all_apps_user()
    client.force_login(user)

    chat = Chat.objects.create(user=user)
    msg = Message.objects.create(chat=chat, text="", is_bot=True)
    sf = SavedFile.objects.create(content_type="text/html")
    sf.file.save("page.html", ContentFile(b"<html><body>Hello</body></html>"))
    cf = ChatFile.objects.create(message=msg, filename="page.html", saved_file=sf)

    url = reverse("chat_next:preview_file", args=[cf.id])
    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode()
    assert 'data-copyable="true"' in content
    assert "preview-html-source" in content
    assert "preview-iframe-html" in content
    assert "srcdoc" in content
    assert "sandbox" in content
