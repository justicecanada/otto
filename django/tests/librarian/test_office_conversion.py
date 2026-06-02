from io import BytesIO

from django.core.files.base import ContentFile
from django.urls import reverse

import pytest

from librarian.models import Document, SavedFile, SavedFileDerivative
from librarian.tasks import (
    _convert_and_replace_legacy_word_document,
    process_document,
)
from librarian.utils.derivatives import DERIVATION_DOC_TO_DOCX
from librarian.utils.office import WORDPROCESSINGML_DOCUMENT_MIME
from librarian.utils.process_engine import extract_markdown, guess_content_type


def test_guess_content_type_recognizes_legacy_word_extension():
    assert guess_content_type(b"legacy-word", path="policy.doc") == "application/msword"


def test_extract_legacy_word_converts_with_libreoffice_bridge(sample_docx, monkeypatch):
    calls = {}

    def fake_convert(content, *, source_filename=None, timeout_seconds=120):
        calls["content"] = content
        calls["source_filename"] = source_filename
        return sample_docx

    monkeypatch.setattr(
        "librarian.utils.process_engine.convert_legacy_word_to_docx",
        fake_convert,
    )

    extraction_result = extract_markdown(
        b"legacy-binary-doc-content",
        "WORD_LEGACY",
        content_type="application/msword",
    )

    assert calls["content"] == b"legacy-binary-doc-content"
    assert "Test Heading" in extraction_result.markdown
    assert "Test paragraph" in extraction_result.markdown
    assert extraction_result.chunks


@pytest.mark.django_db
def test_convert_and_replace_legacy_word_document_updates_saved_file(
    sample_docx, monkeypatch
):
    from docx import Document as DocxDocument

    original_saved_file = SavedFile.objects.create(content_type="application/msword")
    original_saved_file.file.save("legacy.doc", ContentFile(b"legacy doc bytes"))
    original_saved_file.generate_hash()

    document = Document.objects.create(
        saved_file=original_saved_file,
        filename="legacy.doc",
        url="https://example.com/legacy.doc",
        url_content_type="application/msword",
    )

    monkeypatch.setattr(
        "librarian.tasks.convert_legacy_word_to_docx",
        lambda content, *, source_filename=None, timeout_seconds=120: sample_docx,
    )

    converted_content, converted_content_type = (
        _convert_and_replace_legacy_word_document(document, b"legacy doc bytes")
    )

    document.refresh_from_db()

    assert converted_content == sample_docx
    assert converted_content_type == WORDPROCESSINGML_DOCUMENT_MIME
    assert document.filename == "legacy.docx"
    assert document.manual_title == "legacy.doc"
    assert document.original_saved_file_id == original_saved_file.id
    assert document.original_filename == "legacy.doc"
    assert document.url_content_type == WORDPROCESSINGML_DOCUMENT_MIME
    assert document.saved_file_id != original_saved_file.id
    assert document.saved_file.content_type == WORDPROCESSINGML_DOCUMENT_MIME
    assert SavedFileDerivative.objects.filter(
        source_saved_file=original_saved_file,
        derived_saved_file=document.saved_file,
        derivation_type=DERIVATION_DOC_TO_DOCX,
    ).exists()
    with document.saved_file.file.open("rb") as saved_docx:
        persisted_bytes = saved_docx.read()
    assert persisted_bytes == sample_docx
    opened_docx = DocxDocument(BytesIO(persisted_bytes))
    assert opened_docx.paragraphs[0].text == "Test Heading"


@pytest.mark.django_db
def test_download_document_returns_openable_converted_docx(
    client, all_apps_user, sample_docx, monkeypatch
):
    from docx import Document as DocxDocument

    user = all_apps_user()
    client.force_login(user)

    original_saved_file = SavedFile.objects.create(content_type="application/msword")
    original_saved_file.file.save("legacy.doc", ContentFile(b"legacy doc bytes"))
    original_saved_file.generate_hash()

    data_source = user.personal_library.data_sources.create(name="Converted docs")
    document = Document.objects.create(
        data_source=data_source,
        saved_file=original_saved_file,
        filename="legacy.doc",
    )

    monkeypatch.setattr(
        "librarian.tasks.convert_legacy_word_to_docx",
        lambda content, *, source_filename=None, timeout_seconds=120: sample_docx,
    )

    _convert_and_replace_legacy_word_document(document, b"legacy doc bytes")
    document.refresh_from_db()

    response = client.get(
        reverse("librarian:download_document", kwargs={"document_id": document.id})
    )

    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="legacy.docx"'
    downloaded_bytes = b"".join(response.streaming_content)
    assert downloaded_bytes == sample_docx
    opened_docx = DocxDocument(BytesIO(downloaded_bytes))
    assert opened_docx.paragraphs[0].text == "Test Heading"


@pytest.mark.django_db
def test_download_original_document_returns_source_doc(
    client, all_apps_user, sample_docx, monkeypatch
):
    user = all_apps_user()
    client.force_login(user)

    original_bytes = b"legacy doc bytes"
    original_saved_file = SavedFile.objects.create(content_type="application/msword")
    original_saved_file.file.save("legacy.doc", ContentFile(original_bytes))
    original_saved_file.generate_hash()

    data_source = user.personal_library.data_sources.create(name="Converted docs")
    document = Document.objects.create(
        data_source=data_source,
        saved_file=original_saved_file,
        filename="legacy.doc",
    )

    monkeypatch.setattr(
        "librarian.tasks.convert_legacy_word_to_docx",
        lambda content, *, source_filename=None, timeout_seconds=120: sample_docx,
    )

    _convert_and_replace_legacy_word_document(document, original_bytes)
    document.refresh_from_db()

    response = client.get(
        reverse(
            "librarian:download_original_document",
            kwargs={"document_id": document.id},
        )
    )

    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="legacy.doc"'
    assert b"".join(response.streaming_content) == original_bytes

    modal_response = client.get(
        reverse("librarian:modal_view_document", kwargs={"document_id": document.id})
    )
    content = modal_response.content.decode()
    assert (
        reverse("librarian:download_document", kwargs={"document_id": document.id})
        in content
    )
    assert (
        reverse(
            "librarian:download_original_document", kwargs={"document_id": document.id}
        )
        in content
    )
    assert "Processed file" in content
    assert "Original file" in content


@pytest.mark.django_db
def test_convert_and_replace_legacy_word_document_reuses_cached_derivative(
    sample_docx, monkeypatch
):
    first_saved_file = SavedFile.objects.create(content_type="application/msword")
    first_saved_file.file.save("legacy.doc", ContentFile(b"legacy doc bytes"))
    first_saved_file.generate_hash()

    first_document = Document.objects.create(
        saved_file=first_saved_file,
        filename="legacy.doc",
    )

    monkeypatch.setattr(
        "librarian.tasks.convert_legacy_word_to_docx",
        lambda content, *, source_filename=None, timeout_seconds=120: sample_docx,
    )
    _convert_and_replace_legacy_word_document(first_document, b"legacy doc bytes")
    first_document.refresh_from_db()

    second_document = Document.objects.create(
        saved_file=first_saved_file,
        filename="renamed.doc",
    )

    def should_not_run(*args, **kwargs):
        raise AssertionError(
            "LibreOffice conversion should reuse the cached derivative"
        )

    monkeypatch.setattr(
        "librarian.tasks.convert_legacy_word_to_docx",
        should_not_run,
    )

    converted_content, converted_content_type = (
        _convert_and_replace_legacy_word_document(
            second_document,
            b"legacy doc bytes",
        )
    )
    second_document.refresh_from_db()

    assert converted_content == sample_docx
    assert converted_content_type == WORDPROCESSINGML_DOCUMENT_MIME
    assert second_document.saved_file_id == first_document.saved_file_id
    assert second_document.original_saved_file_id == first_saved_file.id
    assert second_document.original_filename == "renamed.doc"
    assert second_document.filename == "renamed.docx"


@pytest.mark.django_db
def test_process_document_reports_legacy_word_conversion_status(
    sample_docx, monkeypatch
):
    saved_file = SavedFile.objects.create(content_type="application/msword")
    saved_file.file.save("legacy.doc", ContentFile(b"legacy doc bytes"))
    saved_file.generate_hash()

    document = Document.objects.create(
        saved_file=saved_file,
        filename="legacy.doc",
        url_content_type="application/msword",
    )

    status_updates = []

    class FakeCurrentTask:
        def update_state(self, state, meta):
            status_updates.append((state, meta))

    class DummyAsyncResult:
        id = "finalize-task-id"
        backend = None

    class DummyExtractionResult:
        needs_azure = False
        chunks = ["chunk"]

    monkeypatch.setattr("librarian.tasks.current_task", FakeCurrentTask())
    monkeypatch.setattr(
        "librarian.tasks._bind_librarian_context", lambda **kwargs: None
    )
    monkeypatch.setattr("librarian.tasks.check_cancel", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        "librarian.tasks._fetch_and_detect",
        lambda doc, refresh_from_url=False: (
            b"legacy doc bytes",
            "application/msword",
            None,
        ),
    )
    monkeypatch.setattr(
        "librarian.tasks._convert_and_replace_legacy_word_document",
        lambda doc, content: (sample_docx, WORDPROCESSINGML_DOCUMENT_MIME),
    )
    monkeypatch.setattr(
        "librarian.tasks._is_zip_container", lambda *args, **kwargs: False
    )
    monkeypatch.setattr(
        "librarian.tasks._extract_and_persist",
        lambda *args, **kwargs: DummyExtractionResult(),
    )
    monkeypatch.setattr(
        "librarian.tasks._should_pause_large_document_embedding",
        lambda *args, **kwargs: False,
    )
    monkeypatch.setattr(
        "librarian.tasks.set_celery_task_id", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "librarian.tasks.finalize_document_light.apply_async",
        lambda **kwargs: DummyAsyncResult(),
    )

    result = process_document.run(document.id, language="en")

    assert result["ok"] is True
    assert (
        "PROCESSING",
        {"status_text": "Converting .doc to .docx..."},
    ) in status_updates
